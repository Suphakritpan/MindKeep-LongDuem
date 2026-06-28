"""Per-type text extraction (docs/SCOPE.md §3.3).

Heavy libraries are imported lazily inside each function so this module loads
even when an optional dependency (or the tesseract binary) is missing.
"""
import csv
import io
from pathlib import Path

MAX_CHARS = 20_000  # cap stored extracted text
_MAX_ROWS = 100


def extract_text(data: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        text = _pdf(data)
    elif ext == ".docx":
        text = _docx(data)
    elif ext == ".xlsx":
        text = _xlsx(data)
    elif ext == ".csv":
        text = _csv(data)
    elif ext in {".jpg", ".jpeg", ".png", ".webp"}:
        text = _ocr(data)
    else:
        text = ""
    return text[:MAX_CHARS].strip()


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs)


def _xlsx(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        lines.append(f"# Sheet: {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= _MAX_ROWS:
                break
            lines.append(", ".join("" if c is None else str(c) for c in row))
    return "\n".join(lines)


def _csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    lines: list[str] = []
    for i, row in enumerate(csv.reader(io.StringIO(text))):
        if i >= _MAX_ROWS:
            break
        lines.append(", ".join(row))
    return "\n".join(lines)


def _ocr(data: bytes) -> str:
    # OCR is best-effort: if pytesseract / the tesseract binary is unavailable,
    # return empty text — the image metadata is still stored (SCOPE §3.3).
    try:
        import pytesseract
        from PIL import Image

        return pytesseract.image_to_string(Image.open(io.BytesIO(data)))
    except Exception:  # noqa: BLE001 — any OCR failure degrades to empty text
        return ""
